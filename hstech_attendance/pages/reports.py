# -*- coding: utf-8 -*-
import calendar
from datetime import datetime
from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QComboBox,
                              QPushButton, QDateEdit, QTableWidget, QTableWidgetItem,
                              QHeaderView, QFrame, QFileDialog, QMessageBox)
from PyQt5.QtCore import Qt, QDate
from PyQt5.QtGui import QFont

import db
from pages.add_employee import DEPARTMENTS

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False


class ReportsPage(QWidget):
    def __init__(self, fixed_ma_nv=None, personal_title=None):
        super().__init__()
        # fixed_ma_nv: nếu được truyền vào (tài khoản nhân viên), trang chỉ
        # hiển thị & xuất báo cáo của đúng nhân viên đó — không thấy số liệu
        # của người khác, không có bộ lọc phòng ban.
        self.fixed_ma_nv = fixed_ma_nv
        self.personal_title = personal_title
        super_title = personal_title or "Báo cáo chấm công"
        self._build_ui(super_title)
        self.refresh()

    def _build_ui(self, page_title):
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 20, 24, 20)
        root.setSpacing(14)

        header = QHBoxLayout()
        title = QLabel(page_title)
        title.setFont(QFont("Segoe UI", 15, QFont.Bold))
        header.addWidget(title)
        header.addStretch()

        self.month_edit = QDateEdit(calendarPopup=True)
        self.month_edit.setDisplayFormat("MM/yyyy")
        self.month_edit.setDate(QDate.currentDate())
        header.addWidget(self.month_edit)

        if not self.fixed_ma_nv:
            self.dept_combo = QComboBox()
            self.dept_combo.addItems(["Tất cả"] + DEPARTMENTS)
            header.addWidget(self.dept_combo)
        else:
            self.dept_combo = None

        export_btn = QPushButton("⬇ Xuất Excel")
        export_btn.setObjectName("primaryBtn")
        export_btn.clicked.connect(self.export_excel)
        header.addWidget(export_btn)

        refresh_btn = QPushButton("Lọc")
        refresh_btn.setObjectName("ghostBtn")
        refresh_btn.clicked.connect(self.refresh)
        header.addWidget(refresh_btn)
        root.addLayout(header)

        # Thẻ tổng hợp số liệu tháng
        summary_row = QHBoxLayout()
        self.summary_labels = {}
        for key, label, color in [("so_ngay_cong", "Số ngày công", "#16A34A"),
                                   ("di_muon", "Đi muộn", "#F59E0B"),
                                   ("ve_som", "Về sớm", "#2563EB"),
                                   ("vang", "Vắng", "#DC2626")]:
            card = QFrame()
            card.setObjectName("card")
            card_l = QVBoxLayout(card)
            val = QLabel("0")
            val.setFont(QFont("Segoe UI", 20, QFont.Bold))
            val.setStyleSheet(f"color:{color};")
            val.setAlignment(Qt.AlignCenter)
            lbl = QLabel(label)
            lbl.setAlignment(Qt.AlignCenter)
            lbl.setStyleSheet("color:#64748B;")
            card_l.addWidget(val)
            card_l.addWidget(lbl)
            summary_row.addWidget(card)
            self.summary_labels[key] = val
        root.addLayout(summary_row)

        if not self.fixed_ma_nv:
            # Chế độ admin: bảng tổng hợp theo từng nhân viên
            self.table = QTableWidget(0, 5)
            self.table.setHorizontalHeaderLabels(
                ["Nhân viên", "Số ngày công", "Đi muộn", "Về sớm", "Vắng"])
            self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        else:
            # Chế độ cá nhân: bảng chi tiết từng ngày trong tháng
            self.table = QTableWidget(0, 4)
            self.table.setHorizontalHeaderLabels(["Ngày", "Giờ vào", "Giờ ra", "Trạng thái"])
            self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        root.addWidget(self.table, 1)

    def refresh(self):
        month = self.month_edit.date().month()
        year = self.month_edit.date().year()
        dept = self.dept_combo.currentText() if self.dept_combo else "Tất cả"
        self._report_cache = db.get_monthly_report(month, year, dept, ma_nv=self.fixed_ma_nv)

        totals = {"so_ngay_cong": 0, "di_muon": 0, "ve_som": 0, "vang": 0}
        for row in self._report_cache:
            for k in totals:
                totals[k] += row[k]
        for k, lbl in self.summary_labels.items():
            lbl.setText(str(totals[k]))

        self.table.setRowCount(0)
        if not self.fixed_ma_nv:
            for row in self._report_cache:
                r = self.table.rowCount()
                self.table.insertRow(r)
                self.table.setItem(r, 0, QTableWidgetItem(row["ho_ten"]))
                self.table.setItem(r, 1, QTableWidgetItem(str(row["so_ngay_cong"])))
                self.table.setItem(r, 2, QTableWidgetItem(str(row["di_muon"])))
                self.table.setItem(r, 3, QTableWidgetItem(str(row["ve_som"])))
                self.table.setItem(r, 4, QTableWidgetItem(str(row["vang"])))
        else:
            last_day = calendar.monthrange(year, month)[1]
            tu_ngay = datetime(year, month, 1)
            den_ngay = datetime(year, month, last_day)
            rows = db.get_history(tu_ngay=tu_ngay, den_ngay=den_ngay, ma_nv=self.fixed_ma_nv)
            rows.sort(key=lambda r: datetime.strptime(r["ngay"], "%d/%m/%Y"))
            for row in rows:
                r = self.table.rowCount()
                self.table.insertRow(r)
                self.table.setItem(r, 0, QTableWidgetItem(row["ngay"]))
                self.table.setItem(r, 1, QTableWidgetItem(row["gio_vao"] or "-"))
                self.table.setItem(r, 2, QTableWidgetItem(row["gio_ra"] or "-"))
                self.table.setItem(r, 3, QTableWidgetItem(row["trang_thai"] or ""))

    def export_excel(self):
        if not OPENPYXL_OK:
            QMessageBox.warning(self, "Thiếu thư viện", "Cần cài đặt openpyxl: pip install openpyxl")
            return
        default_name = (f"ChamCong_{self.fixed_ma_nv}_{self.month_edit.date().toString('MM_yyyy')}.xlsx"
                         if self.fixed_ma_nv else
                         f"BaoCao_ChamCong_{self.month_edit.date().toString('MM_yyyy')}.xlsx")
        path, _ = QFileDialog.getSaveFileName(self, "Xuất báo cáo Excel", default_name, "Excel Files (*.xlsx)")
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Chấm công"
        header_fill = PatternFill(start_color="0B1E3F", end_color="0B1E3F", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D0D0D0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        if not self.fixed_ma_nv:
            ws.append(["Nhân viên", "Mã NV", "Số ngày công", "Đi muộn", "Về sớm", "Vắng"])
            for row in self._report_cache:
                ws.append([row["ho_ten"], row["ma_nv"], row["so_ngay_cong"],
                           row["di_muon"], row["ve_som"], row["vang"]])
            widths = [28, 12, 14, 10, 10, 8]
        else:
            ws.append(["Ngày", "Giờ vào", "Giờ ra", "Trạng thái"])
            for r in range(self.table.rowCount()):
                ws.append([self.table.item(r, c).text() if self.table.item(r, c) else ""
                           for c in range(4)])
            widths = [16, 12, 12, 14]

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
        for r in ws.iter_rows(min_row=2):
            for cell in r:
                cell.border = border
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w

        wb.save(path)
        QMessageBox.information(self, "Thành công", f"Đã xuất báo cáo:\n{path}")
